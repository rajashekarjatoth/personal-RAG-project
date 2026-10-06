"""
Document Loaders module for OmniRAG Studio.
Handles ingestion of diverse file types: PDF, DOCX, TXT, MD, and raw text in-memory.
"""

from typing import List, Dict, Any, Union, BinaryIO, Optional
import os
import io
import pypdf
import docx


class Document:
    """Represents an ingested document with content and rich metadata."""
    def __init__(self, content: str, metadata: Optional[Dict[str, Any]] = None):
        self.content = content.strip()
        self.metadata = metadata or {}
        
    @property
    def word_count(self) -> int:
        return len(self.content.split())
        
    @property
    def char_count(self) -> int:
        return len(self.content)
        
    def to_dict(self) -> Dict[str, Any]:
        return {
            "content": self.content,
            "word_count": self.word_count,
            "char_count": self.char_count,
            "metadata": self.metadata
        }


class DocumentLoader:
    """Unified loader providing loaders for PDF, DOCX, TXT, Markdown, and raw strings."""
    
    @staticmethod
    def load_from_text(text: str, title: str = "Raw Document", metadata: Optional[Dict[str, Any]] = None) -> Document:
        """Loads a document from a raw string."""
        meta = metadata or {}
        meta.setdefault("title", title)
        meta.setdefault("file_type", "raw_text")
        return Document(content=text, metadata=meta)

    @staticmethod
    def load_from_txt_file(file_path_or_buffer: Union[str, BinaryIO], filename: str = "document.txt") -> Document:
        """Loads a document from a TXT or Markdown file or buffer."""
        if isinstance(file_path_or_buffer, str):
            with open(file_path_or_buffer, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
            name = os.path.basename(file_path_or_buffer)
        else:
            raw = file_path_or_buffer.read()
            content = raw.decode("utf-8", errors="replace") if isinstance(raw, bytes) else str(raw)
            name = filename
            
        return Document(
            content=content,
            metadata={
                "title": name,
                "file_type": "txt",
                "source": name
            }
        )

    @staticmethod
    def load_from_pdf(file_path_or_buffer: Union[str, BinaryIO], filename: str = "document.pdf") -> Document:
        """Extracts text and page metadata from a PDF file or binary stream using pypdf."""
        if isinstance(file_path_or_buffer, str):
            reader = pypdf.PdfReader(file_path_or_buffer)
            name = os.path.basename(file_path_or_buffer)
        else:
            reader = pypdf.PdfReader(file_path_or_buffer)
            name = filename
            
        extracted_pages = []
        for page_idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                extracted_pages.append(text.strip())
                
        full_text = "\n\n".join(extracted_pages)
        return Document(
            content=full_text,
            metadata={
                "title": name,
                "file_type": "pdf",
                "total_pages": len(reader.pages),
                "source": name
            }
        )

    @staticmethod
    def load_from_docx(file_path_or_buffer: Union[str, BinaryIO], filename: str = "document.docx") -> Document:
        """Extracts text and paragraph count from a Microsoft Word DOCX file or stream."""
        if isinstance(file_path_or_buffer, str):
            doc = docx.Document(file_path_or_buffer)
            name = os.path.basename(file_path_or_buffer)
        else:
            doc = docx.Document(file_path_or_buffer)
            name = filename
            
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        full_text = "\n\n".join(paragraphs)
        
        return Document(
            content=full_text,
            metadata={
                "title": name,
                "file_type": "docx",
                "total_paragraphs": len(paragraphs),
                "source": name
            }
        )

    @classmethod
    def auto_load(cls, file_source: Union[str, BinaryIO], filename: Optional[str] = None) -> Document:
        """Automatically detects format by extension or filename and extracts text."""
        effective_name = filename or (os.path.basename(file_source) if isinstance(file_source, str) else "unknown.txt")
        ext = effective_name.lower().split(".")[-1] if "." in effective_name else ""
        
        if ext == "pdf":
            return cls.load_from_pdf(file_source, filename=effective_name)
        elif ext in ("docx", "doc"):
            return cls.load_from_docx(file_source, filename=effective_name)
        elif ext in ("txt", "md", "markdown", "csv", "json"):
            return cls.load_from_txt_file(file_source, filename=effective_name)
        else:
            # Fallback to plain text reading
            return cls.load_from_txt_file(file_source, filename=effective_name)
